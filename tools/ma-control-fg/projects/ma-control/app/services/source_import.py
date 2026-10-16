"""Read exported SAP CSV/XLSX files without executing Excel macros."""
import csv
import io
import math
import re
import unicodedata
from pathlib import Path

ALIASES = {
    'ob_number': ['XBLNR', 'LIFEX', 'OB', 'Outbound Delivery', 'Reference', 'Reference Document', 'External Delivery'],
    'material_code': ['MATNR', 'Material', 'Material Code', 'Mã hàng', 'Mã vật tư'],
    'batch': ['CHARG', 'Batch', 'Lô'],
    'quantity': ['MENGE', 'LFIMG', 'Quantity', 'Delivery Quantity', 'Unrestricted', 'Unrestricted Use', 'SL', 'Số lượng', 'unrestricted_qty'],
    'unit': ['MEINS', 'VRKME', 'Unit', 'Base Unit of Measure', 'BUn', 'SU', 'ĐVT'],
    'movement_type': ['BWART', 'Movement Type', 'MvT', 'Loại di chuyển'],
    'doc_no': ['MBLNR', 'Material Document', 'Mat. Doc.', 'Số chứng từ'],
    'ib_number': ['VBELN', 'Inbound Delivery', 'Delivery', 'IB'],
    'plant': ['WERKS', 'Plant', 'Plnt'],
    'sloc': ['LGORT', 'Storage Location', 'SLoc'],
    'gr_status': ['WBSTK', 'WBSTA', 'GR Status'],
}
REQUIRED = {'mb51': {'ob_number', 'material_code', 'batch', 'quantity', 'unit', 'movement_type', 'doc_no'},
            'vl06i': {'ob_number', 'material_code', 'batch', 'quantity', 'unit', 'ib_number', 'plant', 'sloc', 'gr_status'},
            'mb52': {'material_code', 'batch', 'quantity', 'unit', 'plant', 'sloc'}}


def _normal(value):
    value = unicodedata.normalize('NFD', str(value or '').lower().replace('đ', 'd'))
    return re.sub('[^a-z0-9]', '', ''.join(c for c in value if not unicodedata.combining(c)))


def _number(value):
    if isinstance(value, bool):
        raise ValueError('Số lượng không được là true/false')
    if isinstance(value, str):
        value = value.strip().replace(' ', '')
        negative = value.endswith('-')
        if negative:
            value = '-' + value[:-1]
        if ',' in value and '.' in value:
            if value.rfind(',') > value.rfind('.'):
                value = value.replace('.', '').replace(',', '.')
            else:
                value = value.replace(',', '')
        elif ',' in value:
            # Ambiguous grouping is rejected; use an unambiguous decimal/export.
            if re.fullmatch(r'-?\d{1,3}(,\d{3})+', value):
                raise ValueError('Số lượng có dấu phẩy phân nhóm mơ hồ; xuất lại dạng số hoặc dấu chấm thập phân')
            value = value.replace(',', '.')
    result = float(value)
    if not math.isfinite(result):
        raise ValueError('Số lượng không hữu hạn')
    return result


def parse_source(kind, body, filename):
    if kind not in REQUIRED:
        raise ValueError('Loại nguồn không hợp lệ')
    if len(body) > 10 * 1024 * 1024:
        raise ValueError('Tệp nguồn vượt 10 MB; chia bộ dữ liệu pilot nhỏ hơn')
    suffix = Path(filename).suffix.lower()
    if suffix in {'.xlsx', '.xlsm'}:
        import zipfile
        with zipfile.ZipFile(io.BytesIO(body)) as package:
            if sum(i.file_size for i in package.infolist()) > 80 * 1024 * 1024:
                raise ValueError('Workbook nguồn quá lớn')
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
        sheets = []
        for sheet in book:
            rows = []
            for row in sheet.iter_rows():
                values = []
                for cell in row:
                    value = cell.value
                    if isinstance(value, (int, float)) and float(value).is_integer() and re.fullmatch('0+', str(cell.number_format)):
                        value = str(int(value)).zfill(len(cell.number_format))
                    values.append(value)
                rows.append(values)
            sheets.append(rows)
        book.close()
    elif suffix in {'.csv', '.tsv', '.txt'}:
        try:
            text = body.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = body.decode('utf-16')
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel_tab if '\t' in text else csv.excel
        sheets = [list(csv.reader(io.StringIO(text), dialect))]
    else:
        raise ValueError('Dùng CSV/TSV/XLSX/XLSM xuất từ SAP; không chạy macro nguồn')
    matches = []
    for rows in sheets:
        for index, row in enumerate(rows[:40]):
            header = {_normal(v): n for n, v in enumerate(row) if v is not None}
            mapping = {}
            for key, aliases in ALIASES.items():
                for alias in [key, *aliases]:
                    if _normal(alias) in header:
                        mapping[key] = header[_normal(alias)]
                        break
            if REQUIRED[kind] <= mapping.keys():
                matches.append((rows, index, mapping))
                break
    if len(matches) != 1:
        raise ValueError('Không tìm được đúng một bảng có đủ cột nguồn; cần file xuất riêng cho mỗi báo cáo')
    rows, start, mapping = matches[0]
    output = []
    for row in rows[start + 1:]:
        if not any(str(v or '').strip() for v in row):
            continue
        values = {k: row[n] if n < len(row) else None for k, n in mapping.items()}
        if any(values[k] is None or str(values[k]).strip() == '' for k in REQUIRED[kind]):
            raise ValueError('Nguồn có dòng tổng/thiếu cột bắt buộc; xuất dữ liệu chi tiết không kèm tổng')
        entry = {k: str(v).strip() for k, v in values.items() if v is not None}
        entry['quantity'] = _number(values['quantity'])
        if kind == 'mb52':
            entry['unrestricted_qty'] = entry.pop('quantity')
            entry['stock_type'] = 'UNRESTRICTED'
        output.append(entry)
    if not output:
        raise ValueError('File nguồn không có dữ liệu chi tiết')
    return output
