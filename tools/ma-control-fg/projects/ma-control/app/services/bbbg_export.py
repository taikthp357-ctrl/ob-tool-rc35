"""Local draft exporter using the approved historical BBBG layout.

This does not execute VBA, SAP, email or printing. CF is not SAP reconciliation.
"""
from copy import copy
from io import BytesIO
from pathlib import Path
import hashlib
import os
import tempfile
import math
from datetime import datetime, timezone

PROJECT_ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = PROJECT_ROOT / 'baseline/fg/extracted/261008_BBBG_FG_V2_2_14.xlsm'


def sha256_file(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def export_bbbg(path, header, items, template=TEMPLATE):
    # Import lazily: missing exporter dependency must fail without a DB success.
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell
    from openpyxl.drawing.image import Image

    if not items:
        raise ValueError('BBBG không có dòng hàng')
    for item in items:
        qty = item.get('cf_actual_qty')
        if qty is None or not math.isfinite(float(qty)) or float(qty) <= 0:
            raise ValueError('BBBG cần số lượng CF thực tế dương cho mọi dòng')
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        raise FileExistsError('Tệp BBBG đã tồn tại; không ghi đè')
    workbook = load_workbook(template, keep_vba=False, keep_links=False)
    source = workbook['BBBG_NEW']
    sheet = workbook.copy_worksheet(source)
    sheet.title = 'BBBG_DRAFT'
    # copy_worksheet deliberately excludes drawings. Keep independent image streams.
    for original_image in source._images:
        logo = Image(BytesIO(original_image._data()))
        logo.anchor = copy(original_image.anchor)
        logo.width, logo.height = original_image.width, original_image.height
        sheet.add_image(logo)
    for other in list(workbook):
        if other != sheet:
            workbook.remove(other)
    workbook.defined_names.clear()
    sheet.defined_names.clear()
    workbook.properties.creator = 'MA Control'
    workbook.properties.lastModifiedBy = 'MA Control'
    workbook.properties.lastPrinted = None
    workbook.properties.created = workbook.properties.modified = datetime.now(timezone.utc).replace(tzinfo=None)
    sheet.sheet_state = 'visible'
    # Move the five footer rows for more than the template's 18 detail slots.
    display_rows = max(18, len(items))
    footer = 13 + display_rows
    footer_cells = [(c.row, c.column, c.value, copy(c._style), copy(c.alignment))
                    for row in sheet.iter_rows(min_row=31, max_row=35, max_col=11)
                    for c in row if not isinstance(c, MergedCell)]
    footer_merges = [str(r) for r in sheet.merged_cells.ranges if r.min_row >= 31]
    for merged in footer_merges:
        sheet.unmerge_cells(merged)
    for row in sheet.iter_rows(min_row=13, max_row=max(35, footer + 4), max_col=11):
        for cell in row:
            if not isinstance(cell, MergedCell):
                cell.value = None
    exemplar = source[13]
    for row_number in range(13, footer):
        sheet.row_dimensions[row_number].height = source.row_dimensions[13].height
        for col in range(1, 12):
            sheet.cell(row_number, col)._style = copy(exemplar[col-1]._style)
    shift = footer - 31
    for row, col, value, style, alignment in footer_cells:
        cell = sheet.cell(row + shift, col, value)
        cell._style, cell.alignment = style, alignment
        sheet.row_dimensions[row + shift].height = source.row_dimensions[row].height
    for merged in footer_merges:
        from openpyxl.worksheet.cell_range import CellRange
        cr = CellRange(merged)
        cr.shift(row_shift=shift)
        sheet.merge_cells(str(cr))
    verified_import = bool(header.get('reconciliation'))
    sheet['A1'] = ('ĐÃ ĐỐI CHIẾU TỆP NHẬP MB51 / VL06I — CHƯA XÁC NHẬN GR/ĐÍNH KÈM SAP'
                   if verified_import else 'BẢN NHÁP — CHƯA ĐỐI CHIẾU MB51 / VL06I; KHÔNG XÁC NHẬN SAP/GR')
    sheet['I1'] = header['slip_no']
    sheet['E1'] = header.get('storage_bin') or ''
    sheet['A4'] = 'Ngày (Date): ' + header['handover_date']
    sheet['C5'], sheet['G5'] = header['sender_name'], header.get('receiver_name') or ''
    sheet['C7'], sheet['G7'] = 'Kho MA', header.get('shop') or ''
    for row_number, item in enumerate(items, 13):
        values = [row_number - 12, item['material_code'], item.get('batch') or '/', '/',
                  ((item.get('ib_number') or '/') + ' / ' + (item.get('doc_no') or header['doc_no'])),
                  item.get('unit') or 'PCS', float(item['cf_actual_qty']),
                  None, None, None, ('Đối chiếu tệp MB51/VL06I' if verified_import else 'CF thực tế; chưa đối chiếu SAP')]
        for col, value in enumerate(values, 1):
            cell = sheet.cell(row_number, col, value)
            if header.get('template_mode') == 'Set':
                font = copy(cell.font)
                font.sz = header['font_size']
                cell.font = font
            # Identifiers and user input must never become executable formulas.
            if isinstance(value, str):
                cell.data_type = 's'
            if col in (2, 3, 5):
                cell.number_format = '@'
        sheet.cell(row_number, 7).number_format = '#,##0.###'
    sheet.cell(footer, 7, sum(float(i['cf_actual_qty']) for i in items))
    sheet.cell(footer, 7).number_format = '#,##0.###'
    sheet.cell(footer + 4, 1, 'Họ và Tên:\n' + header['sender_name'])
    sheet.cell(footer + 4, 3, 'Họ và Tên:\n')
    sheet.cell(footer + 4, 8, 'Họ và Tên:\n' + (header.get('receiver_name') or ''))
    # Security signature note has many wrapped bilingual lines in the baseline.
    for row_number in (footer + 2, footer + 3):
        sheet.row_dimensions[row_number].height = max(sheet.row_dimensions[row_number].height or 0, 60)
    sheet.print_area = f'A1:K{footer + 4}'
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    fd, temporary = tempfile.mkstemp(suffix='.xlsx', dir=path.parent)
    os.close(fd)
    try:
        workbook.save(temporary)
        verified = load_workbook(temporary, data_only=True)
        if verified['BBBG_DRAFT'].cell(footer, 7).value != sum(float(i['cf_actual_qty']) for i in items):
            raise ValueError('Kiểm tra tổng BBBG sau xuất thất bại')
        verified.close()
        # Hard link publishes without overwriting an existing destination.
        os.link(temporary, path)
    finally:
        workbook.close()
        Path(temporary).unlink(missing_ok=True)
    return {'file_path': str(path.resolve()), 'export_path': str(path.resolve()),
            'file_sha256': sha256_file(path), 'export_status': 'LOCAL_SAVED',
            'baseline_sha256': sha256_file(template),
            'is_draft': not verified_import,
            'reconciliation_status': 'IMPORTED_RECONCILED' if verified_import else 'NOT_VERIFIED',
            'source_hash': header.get('reconciliation', {}).get('source_hash')}
