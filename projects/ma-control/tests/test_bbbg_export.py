import sqlite3
from pathlib import Path
import pytest
import sys
import types
from openpyxl import load_workbook
from app.services.bbbg_export import export_bbbg
from app.services.bbbg_service import BBBGService
from app.core.database import init_db


HEADER = dict(slip_no='TEST-001', doc_no='0000123', sender_name='Nguyễn Tài',
              receiver_name='Người nhận', shop='BODY', storage_bin='A01', handover_date='2026-10-13')


def test_template_real_file_actual_quantity_and_long_detail(tmp_path):
    items = [dict(material_code='00001', batch='00002', cf_actual_qty=2.5)] * 24
    result = export_bbbg(tmp_path / 'draft.xlsx', HEADER, items)
    assert result['is_draft'] and result['reconciliation_status'] == 'NOT_VERIFIED'
    w = load_workbook(result['file_path'], data_only=True)
    s = w['BBBG_DRAFT']
    assert s['B13'].value == '00001'
    assert s['G13'].value == 2.5 and s['G37'].value == 60
    assert s['H13'].value is None and s['I13'].value is None
    assert 'MB51' in s['A1'].value
    assert s['A37'].value == 'Tổng (Total)'
    assert s['C5'].value == HEADER['sender_name']
    assert not any(c.data_type == 'f' for row in s for c in row)
    assert s['B13'].style_id == s['B14'].style_id
    assert s.sheet_state == 'visible'
    assert len(s._images) == 1
    assert len(w.defined_names) == 0
    assert w.properties.creator == w.properties.lastModifiedBy == 'MA Control'
    w.close()


@pytest.mark.parametrize('qty', [0, -1, None, float('nan'), float('inf')])
def test_invalid_actual_quantity_no_file(tmp_path, qty):
    with pytest.raises(ValueError):
        export_bbbg(tmp_path / 'draft.xlsx', HEADER, [dict(material_code='M', cf_actual_qty=qty)])
    assert not (tmp_path / 'draft.xlsx').exists()


@pytest.fixture
def database(tmp_path, monkeypatch):
    import app.core.database as database_module
    # Isolate exporter contract from the independently tested SAP reconciliation implementation.
    monkeypatch.setitem(sys.modules, 'app.services.reconciliation_service', types.SimpleNamespace(get_verified_reconciliation=lambda conn, pick_id: None))
    monkeypatch.setattr(database_module, 'DB_PATH', tmp_path / 'isolated.db')
    init_db()
    conn = database_module.get_connection()
    conn.execute('DELETE FROM bbbg_records')
    conn.execute('DELETE FROM pick_order_items')
    conn.execute('DELETE FROM pick_orders')
    conn.execute("DELETE FROM sqlite_sequence WHERE name = 'pick_orders'")
    conn.execute("INSERT INTO pick_orders (pick_order_code, shop, model, status, total_qty) VALUES ('P-TEST', 'BODY', 'VF8', 'USER_CONFIRMED', 10)")
    conn.execute("INSERT INTO pick_order_items (pick_order_id, material_code, batch, qty, cf_actual_qty, cf_status) VALUES (1, 'BIW71001001AA', '001', 10, 3, 'CONFIRMED')")
    conn.commit()
    yield conn
    conn.close()


def test_service_retry_real_evidence_and_no_external_success(database, tmp_path):
    svc = BBBGService(database, tmp_path / 'exports')
    first = svc.create_bbbg(1, 'Tài')
    retry = svc.create_bbbg(1, 'Tài')
    assert first['id'] == retry['id'] and retry['idempotent']
    assert first['total_qty'] == 3
    assert first['gos_bbbg_attached'] == first['general_log_saved'] == first['shop_log_saved'] == 0
    row = database.execute('SELECT * FROM bbbg_records').fetchone()
    assert row['unc_export_status'] == 'PENDING' and row['status'] == 'DRAFT_CREATED'
    assert database.execute('SELECT COUNT(*) FROM bbbg_records').fetchone()[0] == 1
    assert Path(svc.get_export_path(first['id'])).is_file()
    Path(first['file_path']).write_bytes(b'changed')
    with pytest.raises(ValueError, match='Tệp BBBG'):
        svc.create_bbbg(1, 'Tài')


def test_export_failure_rolls_back(database, tmp_path, monkeypatch):
    import app.services.bbbg_service as module
    def fail(*args, **kwargs):
        raise OSError('disk full')
    monkeypatch.setattr(module, 'export_bbbg', fail)
    with pytest.raises(OSError):
        BBBGService(database, tmp_path / 'exports').create_bbbg(1, 'Tài')
    assert database.execute('SELECT COUNT(*) FROM bbbg_records').fetchone()[0] == 0
    assert database.execute('SELECT status FROM pick_orders WHERE id=1').fetchone()[0] == 'USER_CONFIRMED'


def test_draft_then_verified_upgrade_same_record(database, tmp_path, monkeypatch):
    svc = BBBGService(database, tmp_path / 'exports')
    first = svc.create_bbbg(1, 'Tài')
    source = {'source_hash': 'test-evidence-hash', 'rows': [{'material_code': 'BIW71001001AA', 'batch': '001', 'quantity': 3, 'unit': 'EA', 'ib_number': '000011', 'doc_no': '000022'}]}
    monkeypatch.setitem(sys.modules, 'app.services.reconciliation_service', types.SimpleNamespace(get_verified_reconciliation=lambda conn, pick_id: source))
    with pytest.raises(ValueError, match='người nhận'):
        svc.create_bbbg(1, 'Tài')
    upgraded = svc.create_bbbg(1, 'Tài', receiver_name='Nguyễn An', storage_bin='A02')
    assert upgraded['id'] == first['id'] and upgraded['slip_no'] == first['slip_no']
    assert not upgraded['is_draft'] and upgraded['finalized']
    assert database.execute('SELECT COUNT(*) FROM bbbg_records').fetchone()[0] == 1
    assert database.execute('SELECT status FROM bbbg_records').fetchone()[0] == 'VERIFIED_IMPORTED'
    assert Path(first['file_path']).exists() and Path(upgraded['file_path']).exists()
    assert svc.get_export_path(first['id']) == upgraded['file_path']
    w = load_workbook(upgraded['file_path'], data_only=True)
    assert w['BBBG_DRAFT']['E13'].value == '000011 / 000022'
    assert w['BBBG_DRAFT']['F13'].value == 'EA'
    assert w['BBBG_DRAFT']['G5'].value == 'Nguyễn An'
    frozen = svc.create_bbbg(1, 'Tên khác', receiver_name='Người khác')
    assert frozen['idempotent'] and frozen['receiver_name'] == 'Nguyễn An'
    assert frozen['file_sha256'] == upgraded['file_sha256']


def test_db_failure_after_file_creation_removes_uncommitted_export(database, tmp_path):
    database.execute("CREATE TRIGGER reject_bbbg BEFORE INSERT ON bbbg_records BEGIN SELECT RAISE(ABORT, 'ledger unavailable'); END")
    database.commit()
    output_dir = tmp_path / 'exports'
    with pytest.raises(sqlite3.IntegrityError, match='ledger unavailable'):
        BBBGService(database, output_dir).create_bbbg(1, 'Tài')
    assert database.execute('SELECT COUNT(*) FROM bbbg_records').fetchone()[0] == 0
    assert list(output_dir.glob('*.xlsx')) == []


def test_unknown_model_never_fabricates_vf8(database, tmp_path):
    svc = BBBGService(database, tmp_path / 'exports')
    with pytest.raises(ValueError, match='mapping model'):
        svc.resolve_model('UNMAPPED')
    assert svc.resolve_model('UNMAPPED', 'VF7') == 'VF7'
    with pytest.raises(ValueError, match='mapping model'):
        svc.resolve_model('UNMAPPED', 'UNKNOWN')
